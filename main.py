from docx import Document
from langchain_core.prompts import ChatPromptTemplate
from langchain_community.llms import VLLMOpenAI
from langchain_core.output_parsers import StrOutputParser
from tqdm import tqdm
import time

# Конфигурация
API_URL = "http://192.168.1.87:1234/v1"
LLM_NAME = "google/gemma-2-9b"
API_TOKEN = "ваш_токен"
INPUT_FILE = "input.docx"
OUTPUT_FILE = "translated.docx"

# Создаем цепочку для перевода
def create_translation_chain():
    prompt = ChatPromptTemplate.from_template(
        "Переведи текст с английского на русский. Сохрани числа, спецсимволы и форматирование. "
        "Текст: {text}"
    )
    
    llm = VLLMOpenAI(
        openai_api_key=API_TOKEN,
        openai_api_base=API_URL,
        model_name=LLM_NAME,
        max_tokens=4096,
        temperature=0.1,
        top_p=0.9,
    )
    
    return prompt | llm | StrOutputParser()

# Упрощенный перевод с ретраями
def translate_text(translator, text, max_retries=3):
    if not text.strip():
        return text
        
    for attempt in range(max_retries):
        try:
            return translator.invoke({"text": text})
        except Exception as e:
            print(f"Ошибка перевода ({attempt+1}/{max_retries}): {e}")
            time.sleep(2 ** attempt)  # Экспоненциальная задержка
    
    print(f"Не удалось перевести фрагмент: {text[:100]}...")
    return text  # Возвращаем оригинал при неудаче

# Основная функция обработки документа
def process_document():
    print("Загрузка документа...")
    doc = Document(INPUT_FILE)
    new_doc = Document()
    
    # Создание переводчика
    translator = create_translation_chain()
    
    # Подсчет элементов для прогресс-бара
    total_elements = len(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            total_elements += len(row.cells)
    
    # Обработка элементов документа
    with tqdm(total=total_elements, desc="Перевод документа") as pbar:
        # Обработка параграфов
        for para in doc.paragraphs:
            if para.text.strip():
                # Переводим весь текст параграфа как единый блок
                translated_text = translate_text(translator, para.text)
                new_para = new_doc.add_paragraph(translated_text)
            else:
                new_doc.add_paragraph()  # Пустой параграф
            
            pbar.update(1)
        
        # Обработка таблиц
        for table in doc.tables:
            new_table = new_doc.add_table(rows=len(table.rows), cols=len(table.columns))
            
            for i, row in enumerate(table.rows):
                for j, cell in enumerate(row.cells):
                    if cell.text.strip():
                        # Переводим весь текст ячейки как единый блок
                        translated_text = translate_text(translator, cell.text)
                        new_table.cell(i, j).text = translated_text
                    else:
                        new_table.cell(i, j).text = ""  # Пустая ячейка
                    
                    pbar.update(1)
    
    # Сохранение документа
    new_doc.save(OUTPUT_FILE)
    print(f"\n✅ Перевод завершен! Файл сохранен как: {OUTPUT_FILE}")
    print(f"Исходный документ: {len(doc.paragraphs)} параграфов, {len(doc.tables)} таблиц")
    print(f"Переведенный документ: {len(new_doc.paragraphs)} параграфов, {len(new_doc.tables)} таблиц")

if __name__ == "__main__":
    process_document()