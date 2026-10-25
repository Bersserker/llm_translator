from pathlib import Path

import pytest
from docx import Document

from tech_translator.config import Settings
from tech_translator.core.pipeline import translate_document
from tech_translator.core.segmenter import segment


class UpperTranslator:
    def __init__(self):
        self.calls = []

    def translate(self, text, **kwargs):
        self.calls.append((text, kwargs))
        return text.upper()


@pytest.mark.parametrize("text", ["Hello world. Next sentence!\n\nParagraph.", "a" * 100, "Привет мир 😀" * 8, "   \n"])
def test_segmentation(text):
    chunks = segment(text, 15, max_bytes=20)
    assert "".join(chunks) == text
    assert all(len(chunk) <= 15 and len(chunk.encode("utf-8")) <= 20 for chunk in chunks)


def test_text_structure_and_context(tmp_path):
    source, target = tmp_path / "source.txt", tmp_path / "target.txt"
    source.write_bytes(b" Hello world.\r\n\r\nSecond paragraph.  \r\n")
    translator = UpperTranslator()
    translate_document(source, target, Settings(), translator=translator)
    assert target.read_bytes() == b" HELLO WORLD.\r\n\r\nSECOND PARAGRAPH.  \r\n"
    assert "Hello world." in translator.calls[-1][1]["context"]


def test_markdown_structure(tmp_path):
    source, target = tmp_path / "source.md", tmp_path / "target.md"
    source.write_text("# Heading\n\n- **Bold** and `literal code` [Link](https://host/path(x))\n"
                      "\n```python\nprint('unchanged')\n```\n\n"
                      "    code block\n\n[ref]: https://example.org\n\n"
                      "| Name | Value |\n| --- | --- |\n| server | 12 |\n", encoding="utf-8")
    translate_document(source, target, Settings(), translator=UpperTranslator())
    output = target.read_text()
    assert "# HEADING" in output
    assert "- **BOLD** AND `literal code` [LINK](https://host/path(x))" in output
    assert "```python\nprint('unchanged')\n```" in output
    assert "    code block" in output
    assert "[ref]: https://example.org" in output
    assert "| NAME | VALUE |\n| --- | --- |\n| SERVER | 12 |" in output


def test_docx_structure(tmp_path):
    source, target = tmp_path / "source.docx", tmp_path / "target.docx"
    document = Document()
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.add_run("Heading").bold = True
    document.add_paragraph("Body")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Cell"
    document.sections[0].header.paragraphs[0].text = "Header"
    document.sections[0].footer.paragraphs[0].text = "Footer"
    document.save(source)
    translate_document(source, target, Settings(), translator=UpperTranslator())
    result = Document(target)
    assert result.paragraphs[0].text == "HEADING"
    assert result.paragraphs[0].runs[0].bold
    assert result.paragraphs[0].style.name == "Heading 1"
    assert result.tables[0].cell(0, 0).text == "CELL"
    assert result.sections[0].header.paragraphs[0].text == "HEADER"
    assert result.sections[0].footer.paragraphs[0].text == "FOOTER"


def test_existing_output_and_input_protected(tmp_path):
    source, target = tmp_path / "a.txt", tmp_path / "b.txt"
    source.write_text("Hello")
    target.write_text("Existing")
    with pytest.raises(FileExistsError):
        translate_document(source, target, Settings(), translator=UpperTranslator())
    assert target.read_text() == "Existing"
    with pytest.raises(ValueError):
        translate_document(source, source, Settings(), force=True)
    translate_document(source, target, Settings(), force=True, translator=UpperTranslator())
    assert target.read_text() == "HELLO"


def test_failure_leaves_no_partial_output(tmp_path):
    source, target = tmp_path / "a.txt", tmp_path / "b.txt"
    source.write_text("Hello\nSecond paragraph")
    class Broken:
        def translate(self, *args, **kwargs):
            raise ConnectionError("offline")
    with pytest.raises(ConnectionError):
        translate_document(source, target, Settings(), translator=Broken())
    assert not target.exists()


def test_glossary_passed_to_translator(tmp_path):
    source, target = tmp_path / "a.txt", tmp_path / "b.txt"
    glossary = tmp_path / "glossary.yaml"
    glossary.write_text("server: сервер", encoding="utf-8")
    source.write_text("server")
    translator = UpperTranslator()
    translate_document(source, target, Settings(translation={"glossary_path": glossary}), translator=translator)
    assert translator.calls[0][1]["glossary"] == {"server": "сервер"}


def test_markdown_metadata_and_references(tmp_path):
    source, target = tmp_path / "a.md", tmp_path / "b.md"
    source.write_text("---\ntitle: Original title\n---\n\n- [x] Done\n\n"
                      "Read [reference] and [reference][] and [Label][reference].\n\n"
                      "[reference]: https://example.org\n")
    translate_document(source, target, Settings(), translator=UpperTranslator())
    text = target.read_text()
    assert "title: Original title" in text
    assert "- [x] DONE" in text
    assert "[reference] AND [reference][] AND [LABEL][reference]" in text


def test_utf8_bom(tmp_path):
    source, target = tmp_path / "a.txt", tmp_path / "b.txt"
    source.write_bytes(b"\xef\xbb\xbfHello\r\n")
    translate_document(source, target, Settings(), translator=UpperTranslator())
    assert target.read_bytes() == b"\xef\xbb\xbfHELLO\r\n"
